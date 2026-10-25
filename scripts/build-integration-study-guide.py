"""Short oral-defense guide; run with the Codex bundled Python."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/entregables/sprint-2/Guia_Defensa_Pruebas_Integracion_Sprint_2_Brotar.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.8)
sec.top_margin = sec.bottom_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    pp = style.element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.03
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(9)
    doc.styles[name].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'Guía de defensa de pruebas de integración del Sprint 2'
doc.core_properties.author = 'Equipo Brotar'
doc.core_properties.subject = 'Proyecto de Sistemas II'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('BROTAR  |  Guía de estudio  |  ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=2, new=False):
    para = doc.add_heading(text, level)
    if new: para.paragraph_format.page_break_before = True
    return para

def code(text):
    para = p(text)
    para.paragraph_format.line_spacing = 1
    for run in para.runs:
        run.font.name = 'Consolas'; run.font.size = Pt(9)
    return para

def source(path, needle):
    lines = (ROOT / path).read_text(encoding='utf-8').splitlines()
    line = next(i for i, value in enumerate(lines, 1) if needle in value)
    code(f'{path} · línea {line}')

doc.add_paragraph('Pruebas de integración del Sprint 2', 'Title')
p('Guía breve para explicar el código y demostrar sus resultados', True)
p('Proyecto de Sistemas II · Ricardo, Alison y Santiago\nDEV 1fcecd4 · 7 de octubre de 2026')
p('En la defensa debemos mostrar que los componentes se comunican correctamente y que los datos y permisos se conservan. Esta guía reúne un guion oral, una prueba para abrir y los comandos para repetir el ensayo local.')
h('1 Qué decir al comenzar', 1)
p('“En Brotar usamos pruebas de integración para comprobar que varias partes funcionan juntas. La suite levanta una API NestJS temporal, envía solicitudes HTTP reales y utiliza TypeORM con PostgreSQL en Docker. Verificamos las respuestas y lo guardado en la base. En Sprint 2 probamos el recorrido desde el borrador hasta la revisión, la corrección y la publicación. También probamos entradas inválidas, accesos sin permiso y solicitudes simultáneas. Así detectamos problemas que una prueba aislada de una función no encontraría.”')
h('2 Cómo distinguir las pruebas', 1)
p('Unitaria. Comprueba una regla aislada, por ejemplo que un importe negativo se rechaza. Puede reemplazar dependencias por dobles de prueba.')
p('Integración. Comprueba componentes reales trabajando juntos. Algunas pruebas ejercitan API y base; otras comprueban directamente TypeORM y PostgreSQL. No todas recorren las mismas capas.')
p('Prueba de interfaz de inicio a fin. Añade la interacción del usuario con el navegador. Nuestra suite PostgreSQL no automatiza todos los clics del frontend; el recorrido visual es una comprobación complementaria.')
h('3 Qué se conecta en la prueba de campañas', 1)
code('Solicitud HTTP → controlador y validaciones → caso de uso\n→ repositorio TypeORM → PostgreSQL → respuesta comprobada')
p('Docker aloja PostgreSQL; no ejecuta las aserciones. El test crea su propia API en un puerto libre y conserva la cookie para actuar como creador o administrador. No necesita que dejemos el backend habitual iniciado.')
p('Términos para recordar: una aserción compara un resultado real con el esperado; un fixture es un dato preparado para el ensayo; un rollback revierte una transacción que falla.')

h('4 Qué abrir y demostrar al docente', 1, True)
source('backend/integration/campaign-sections-submission.integration.test.ts', "test('S2 completo")
p('En Visual Studio Code abrir este archivo. Buscar test, luego start, try y finally. Explicar la preparación, las acciones, las comprobaciones y la limpieza. No hace falta memorizar todas sus líneas.')
h('Recorrido principal')
p('1. La preparación registra cuentas ficticias, inicia sesión y asigna roles solo a los fixtures. Arranca AppModule real mediante NestFactory. Esto no concede roles a los compañeros ni resuelve la política definitiva de verificación.')
p('2. Se crea un borrador y se guardan plan, presupuesto, financiamiento, historia y portada. Se consulta de nuevo cada bloque y se compara con lo enviado. Una entrada inválida devuelve 400 y conserva el presupuesto anterior.')
p('3. Se intenta enviar una campaña incompleta: devuelve 422 y no crea una revisión. Después de completarla, el envío correcto devuelve 201 y el estado queda IN_REVIEW.')
p('4. ADMIN solicita cambios. El creador corrige y reenvía; se crea otra ronda y permanecen las observaciones históricas. Una decisión sobre una ronda antigua devuelve 409.')
p('5. ADMIN aprueba. El estado APPROVED sigue privado: la consulta pública devuelve 404. La publicación separada permite la consulta con 200. Reiniciar la API no pierde la publicación ni sus rondas.')
h('Ejemplo de dos solicitudes simultáneas')
code("const sent = await Promise.all([\n  call(`${path}/submission`, creator.cookie, {}),\n  call(`${path}/submission`, creator.cookie, {})\n])\nassert.deepEqual(sent.map(r => r.status).sort(), [201, 409])")
p('Qué decir: “Ambas peticiones intentan enviar la misma campaña. Solo una debe crear la revisión. La otra recibe conflicto. Además consultamos PostgreSQL y comprobamos que existe una sola revisión pendiente. Lo importante es que no haya duplicados, no solo que el botón se deshabilite.”')
h('Otras comprobaciones que puedes señalar')
p('Permisos. Sin sesión se espera 401; un usuario sin rol administrativo recibe 403. Para secciones privadas de otra persona se espera 404, sin exponer su contenido.')
p('Persistencia y rollback. database.integration.test.ts guarda usuario y perfil, reconecta y verifica los datos. Con DB_TEST_RESTART=true reinicia PostgreSQL. También fuerza un error en una transacción y comprueba que el registro revertido no existe.')
source('backend/integration/database.integration.test.ts', "test('PostgreSQL")

h('5 Cómo repetir la suite en local', 1, True)
p('Primero completar la instalación V2 y migraciones del README. Docker y PostgreSQL deben estar disponibles. Ejecutar desde la raíz del repositorio; la configuración debe apuntar al contenedor de esta copia, nunca a producción. La suite escribe datos ficticios y los elimina al terminar.')
code("cd .\\backend\npnpm run check\n$brotarPrevNodeEnv = $env:NODE_ENV\n$env:NODE_ENV = 'test'\n$env:ALLOW_DB_TEST_WRITES = 'true'\n$env:DB_TEST_RESTART = 'false'\ntry {\n  node --env-file=../infra/postgres-v2/.env.backend `\n    --test --test-concurrency=1 dist-test/integration/*.test.js\n} finally {\n  Remove-Item Env:\\ALLOW_DB_TEST_WRITES -ErrorAction SilentlyContinue\n  Remove-Item Env:\\DB_TEST_RESTART -ErrorAction SilentlyContinue\n  $env:NODE_ENV = $brotarPrevNodeEnv\n}")
p('Para ensayar también el reinicio del contenedor, cambiar DB_TEST_RESTART a true antes de ejecutar. Interrumpe temporalmente la conexión a PostgreSQL y exige el proyecto V2 esperado; no hacerlo durante otra demostración. No borrar volúmenes ni restaurar SQL para repetir pruebas.')
h('6 Preguntas probables y respuesta breve', 1)
p('¿Cómo comprueban que pasó? “Con aserciones de códigos HTTP, respuestas y registros en PostgreSQL. Si alguna condición no coincide, el test falla.”')
p('¿Cuántas pruebas hay? “El ensayo registrado del 5 de octubre aprobó 19 integraciones. Contienen muchas aserciones; no son 19 historias. Los 300 casos totales incluyen frontend, backend y regresiones.”')
p('¿Se usa una base simulada? “En esta suite usamos PostgreSQL real. Los datos son ficticios y temporales, pero las operaciones y respuestas son reales.”')
p('¿Qué pasa con los datos de prueba? “finally intenta limpiar los fixtures propios y cerrar la API. Una interrupción abrupta puede dejar restos; se revisan esos registros, no se vacían tablas completas.”')
p('¿Un test aprobado asegura que no hay bugs? “No. Demuestra los escenarios y condiciones comprobados; no garantiza todos los comportamientos posibles ni sustituye la revisión del usuario.”')
p('¿Por qué eso importa para Sprint 2? “El flujo cruza sesiones, roles, campañas, revisiones y persistencia. La integración comprueba que el resultado es coherente entre módulos, incluso ante errores y concurrencia.”')
h('7 Qué mostrar como evidencia', 1)
p('Mostrar el archivo del test, una aserción y el resumen de la terminal con pass, fail y skipped. Distinguir el ensayo histórico del 5 de octubre de una ejecución nueva. Para preparar esta guía se revisó el código; no se volvió a ejecutar la suite.')
p('Referencias: README, secciones de integración y último ensayo; backend/integration/campaign-sections-submission.integration.test.ts; database.integration.test.ts; database-target.ts. Las líneas corresponden al corte DEV indicado.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
