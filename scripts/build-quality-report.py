"""Genera la guía de defensa y verifica sus referencias contra el código local.

Ejecutar con el Python de dependencias de Codex (python-docx disponible).
La revisión visual del DOCX es un paso separado, obligatorio antes de entregar.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/calidad-sprint-1.md'
OUTPUT = ROOT / 'docs/entregables/sprint-1/Informe_Calidad_y_Defensa_Sprint_1_Brotar.docx'


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.78)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    # La plantilla base puede incluir una línea azul heredada bajo Title.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    doc.core_properties.title = 'Calidad de código y defensa técnica del Sprint 1 de Brotar'
    doc.core_properties.author = 'Equipo Brotar'
    doc.core_properties.subject = 'Principios y patrones aplicados al código del Sprint 1'
    doc.core_properties.keywords = 'Brotar Sprint 1 SOLID DRY KISS YAGNI'
    doc.core_properties.comments = ''
    references = 0
    for block in SOURCE.read_text(encoding='utf-8').split('\n\n'):
        block = block.strip()
        if not block:
            continue
        if block.startswith('@code '):
            reference_lines = block.splitlines()
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            for index, line in enumerate(reference_lines):
                relative, symbol = line.removeprefix('@code ').split(' | ', 1)
                source_lines = (ROOT / relative).read_text(encoding='utf-8').splitlines()
                found = next((i for i, text in enumerate(source_lines, 1) if symbol in text), None)
                if found is None:
                    raise ValueError(f'Referencia no encontrada: {relative}: {symbol}')
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = index < len(reference_lines) - 1
                run = paragraph.add_run(f'{relative}:{found}')
                run.font.name, run.font.size = 'Consolas', Pt(9)
                references += 1
        elif block.startswith('### '):
            doc.add_heading(block[4:], level=2)
        elif block.startswith('## '):
            doc.add_heading(block[3:], level=1)
        elif block.startswith('# '):
            doc.add_paragraph(block[2:], style='Title')
        else:
            doc.add_paragraph(block)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Brotar  |  Sprint 1  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f'DOCX creado; {references} referencias verificadas. {OUTPUT}')


if __name__ == '__main__':
    build()
