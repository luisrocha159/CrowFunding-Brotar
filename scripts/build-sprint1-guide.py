"""Genera la guía de verificación desde su fuente Markdown, sin alterar el informe anterior."""
from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/verificacion-sprint-1.md'
OUTPUT = ROOT / 'docs/entregables/sprint-1/Guia_Verificacion_y_Demostracion_Sprint_1_Brotar.docx'


def build():
    text = SOURCE.read_text(encoding='utf-8')
    tasks = [t for t in json.loads((ROOT / 'docs/planificacion-sprints.json').read_text(encoding='utf-8'))['tasks'] if t['sprint'] == 1]
    for task in tasks:
        heading_id = task['id'].replace('S1-', 'S1 ')
        assert re.search(r'^## ' + re.escape(heading_id) + r' ', text, re.M), task['id']
        for story in task['stories']:
            section = re.split(r'^## ', text, flags=re.M)
            matching = next(s for s in section if s.startswith(heading_id + ' '))
            assert story in matching, (task['id'], story)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for name, size in [('Normal', 11), ('Title', 24), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
        if name in ['Title', 'Heading 1', 'Heading 2']:
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    for block in text.split('\n\n'):
        block = block.strip()
        if not block:
            continue
        if block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('## '):
            doc.add_heading(block[3:], 1)
        elif block == 'Demostración paso a paso':
            doc.add_heading(block, 2)
        else:
            lines = block.splitlines() if block.startswith('- ') or re.match(r'\d+\. ', block) else [block]
            for line in lines:
                p = doc.add_paragraph(line)
                if line.startswith('- '):
                    p.paragraph_format.left_indent = Inches(.15)
                p.paragraph_format.widow_control = True
    doc.core_properties.title = 'Verificación y guía de demostración del Sprint 1 de Brotar'
    doc.core_properties.author = 'Equipo Brotar'
    doc.core_properties.subject = 'Veinte tareas, checklist, pruebas, recorridos y pendientes'
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Brotar | Sprint 1 | ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f'{OUTPUT}\n20 tareas y referencias BG verificadas')


if __name__ == '__main__':
    build()
