"""Create teacher-facing Sprint 2 reports from the reviewed DEV cut.

Run with Codex's bundled Python. DOCX rendering and visual QA are separate.
"""
import ast
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/entregables/sprint-2'
OUT.mkdir(parents=True, exist_ok=True)
CUT = '1fcecd487e287e3fe30e9bd1f062d52e4fb54c70'
DATE = '7 de octubre de 2026'
PLAN = json.loads((ROOT / 'docs/planificacion-sprints.json').read_text(encoding='utf-8'))
TASKS = {t['id']: t for t in PLAN['tasks'] if t['sprint'] == 2}
source = ast.parse((ROOT / 'scripts/build-sprint2-pdfs.py').read_text(encoding='utf-8'))
DATA = {node.targets[0].id: ast.literal_eval(node.value) for node in source.body
        if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name)
        and node.targets[0].id in {'DEMOS', 'METHODS'}}
DEMOS, METHODS = DATA['DEMOS'], DATA['METHODS']
NAMES = ['Ricardo Nuñez del Prado Justiniano', 'Santiago Rocha Mejia', 'Alison Pizarro Apaza']

def paragraph(doc, text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    p.add_run(text).bold = bold
    return p

def heading(doc, text, level=1):
    assert re.fullmatch(r'[\w\s]+', text, re.UNICODE), text
    return doc.add_heading(text, level=level)

def new_doc(title, subtitle):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin, s.bottom_margin = Inches(.78), Inches(.72)
    s.left_margin = s.right_margin = Inches(.83)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Calibri'); fonts.set(qn('w:hAnsi'), 'Calibri')
        rpr.append(fonts)
        ppr = style.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.08
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Normal'].paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
    doc.styles['Subtitle'].font.size = Pt(12)
    doc.styles['Subtitle'].font.italic = True
    doc.core_properties.author = 'Equipo Brotar'
    doc.core_properties.title = title
    doc.core_properties.subject = 'Proyecto de Sistemas II Sprint 2'
    doc.core_properties.keywords = 'Brotar, Sprint 2, campañas, calidad, defensa técnica'
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('BROTAR   |   Sprint 2   |   ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    paragraph(doc, 'BROTAR', bold=True)
    paragraph(doc, title, 'Title')
    paragraph(doc, subtitle, 'Subtitle')
    paragraph(doc, 'Proyecto de Sistemas II')
    paragraph(doc, 'Equipo de trabajo: Ricardo, Alison y Santiago')
    paragraph(doc, 'Estudiantes', bold=True)
    for name in NAMES: paragraph(doc, name)
    paragraph(doc, 'Corte documentado: rama DEV, commit 1fcecd4')
    paragraph(doc, 'Fecha de preparación: ' + DATE)
    paragraph(doc, 'Destinatario: docente de Proyecto de Sistemas II')
    paragraph(doc, 'Estado: incremento de campañas implementado y publicado; aceptación formal de los líderes y reglas definitivas del cliente pendientes.')
    paragraph(doc, 'Fuentes: tarea oficial de Fase 3, código y pruebas de DEV, planificación interna, README y registros de verificación del 5 de octubre de 2026.')
    paragraph(doc, 'Este informe distingue el funcionamiento verificado de la aprobación académica y de la validación final del cliente. Los 300 casos automatizados incluyen regresiones de Sprint 1 y no representan 300 historias de usuario.')
    return doc

def page(doc, title):
    p = heading(doc, title)
    p.paragraph_format.page_break_before = True

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = str(text)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                m = OxmlElement('w:' + side); m.set(qn('w:w'), '85'); m.set(qn('w:type'), 'dxa'); margins.append(m)
            pr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'bottom', 'left', 'right']:
                b = OxmlElement('w:' + side); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
            pr.append(borders)
            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E7EDF3' if i == 0 else ('F6F6F6' if i % 2 == 0 else 'FFFFFF')); pr.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.02
                if j == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.bold = i == 0
    paragraph(doc, '')
    return t

def code(doc, text):
    p = paragraph(doc, text)
    p.paragraph_format.line_spacing = 1.02
    for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
    return p

def ref(doc, path, symbol=None):
    f = ROOT / path
    assert f.exists(), path
    lines = f.read_text(encoding='utf-8').splitlines()
    number = next((i for i, line in enumerate(lines, 1) if symbol is not None and symbol in line), None)
    if symbol is not None: assert number is not None, (path, symbol)
    code(doc, f'{path}' + (f' · línea {number} · {symbol}' if symbol else ''))

def save(doc, name):
    path = OUT / name
    doc.save(path)
    # Content checks are separate from pagination/visual QA.
    paragraphs = '\n'.join(p.text for p in doc.paragraphs)
    assert all(n in paragraphs for n in NAMES)
    assert 'Sprint 1' not in doc.core_properties.title
    assert '2026' in paragraphs and '1fcecd4' in paragraphs
    print(path)

def task(doc, ident):
    t, demo = TASKS[ident], DEMOS[ident]
    heading(doc, ident.replace('-', ' ') + ' ' + t['title'], 2)
    paragraph(doc, f"Responsable planificado: {t['owner']}. Historias relacionadas: {', '.join(t['stories'])}.")
    paragraph(doc, 'Qué se debía lograr. ' + t['result'])
    paragraph(doc, 'Implementación. ' + demo['impl'])
    paragraph(doc, 'Criterios funcionales comprobados. ' + t['acceptance'][1])
    paragraph(doc, 'Prueba negativa. ' + demo['negative'])
    paragraph(doc, 'Cómo demostrarlo. ' + demo['steps'][0])
    paragraph(doc, 'Evidencia. ' + demo['proof'])
    for r in demo['refs'][:1]: ref(doc, *r)

def report():
    doc = new_doc('Informe detallado del Sprint 2', 'Implementación del módulo de campañas y evidencia de cierre técnico')
    page(doc, '1 Resumen y alcance del incremento')
    paragraph(doc, 'El equipo integró el flujo de campañas sobre el acceso y la persistencia del Sprint 1. Un creador puede conservar varios borradores, completar su propuesta, enviarla y corregir observaciones. Un administrador consulta el expediente completo, solicita cambios, aprueba o rechaza; la publicación es una acción separada. El visitante consulta exclusivamente las campañas publicadas con datos reales.')
    paragraph(doc, 'El objetivo corresponde a la tarea oficial de Fase 3, Módulo de Campañas, secciones 4 a 10 y 13. Se mantuvieron React, NestJS, TypeORM, PostgreSQL y el repositorio existente. La sección 11 excluye aportes, pagos, checkout, desembolsos, devoluciones, finanzas, KYC/KYB y rendición. Configurar meta o recompensa no equivale a implementar esos módulos.')
    table(doc, ['Responsable', 'Tareas vigentes', 'Cantidad'], [
        ['Santiago', 'S2-02, 03, 04, 05, 10, 11 y 13', '7'],
        ['Ricardo', 'S2-12, 14, 15, 17 y 21', '5'],
        ['Alison', 'S2-16', '1']], [1.15, 4.55, 1.1])
    paragraph(doc, 'Los identificadores mantienen la numeración original; los saltos corresponden a tareas trasladadas al inventario futuro. Las 13 tareas no son todo el backlog: se conservan 64 historias generales y 62 tareas distribuidas entre los sprints. El alcance de una tarea puede cubrir solamente parte de una historia BG.')
    heading(doc, 'Criterio de cierre', 2)
    paragraph(doc, 'Para las doce tareas funcionales se contrastan resultado, validación, persistencia, permisos y demostración. S2-21 reúne integración, pruebas, documentación y entrega: el código y sus evidencias están publicados en DEV, pero la revisión formal del alcance requiere a los líderes. No se presenta la aceptación externa como completada por aprobar pruebas.')
    heading(doc, 'Lectura del informe', 2)
    paragraph(doc, 'Las secciones 3 a 9 detallan los compromisos. Las secciones 10 a 12 presentan las validaciones, el recorrido de demostración, las decisiones provisionales y la defensa técnica. El documento de principios y patrones explica por separado la justificación del diseño.')

    page(doc, '2 Arquitectura integración y persistencia')
    paragraph(doc, 'React/Vite ejecuta la interfaz en 127.0.0.1:5173. Las solicitudes relativas /api pasan por el proxy de Vite a NestJS en 127.0.0.1:3000. Solo PostgreSQL corre en Docker; el contenedor V2 publica 5432 como puerto anfitrión 15433, accesible en loopback. NestJS utiliza el rol limitado brotar_app y la base brotar_db, no el superusuario para atender la API.')
    table(doc, ['Capa', 'Responsabilidad', 'Ejemplo'], [
        ['Interfaz', 'Controles, progreso y retroalimentación', 'CampaignSectionsEditor'],
        ['Hook y cliente', 'Estado, solicitudes y contratos', 'useCampaignSections'],
        ['HTTP', 'DTO, autenticación y autorización', 'SectionsController'],
        ['Aplicación', 'Coordinación mediante puertos', 'CampaignSections'],
        ['Dominio', 'Reglas y faltantes sin conexión', 'validateFunding'],
        ['Infraestructura', 'SQL, transacciones y persistencia', 'TypeormSectionsRepository']], [1.05, 3.55, 2.2])
    paragraph(doc, 'La infraestructura vigente es infra/postgres-v2. El volumen conserva usuarios, campañas y sesiones tras detener o reiniciar el contenedor. Git distribuye código y SQL de línea base; no distribuye el volumen, cuentas locales, .env, contraseñas ni archivos privados. Cada instalación prepara su propia configuración siguiendo el README.')
    paragraph(doc, 'La V2 conserva 59 tablas públicas, 5 vistas, 8 secuencias, 31 enums, 16 funciones propias, 40 triggers y 134 claves foráneas. Se reutilizan las estructuras oficiales. Tres migraciones controlan línea base, estado REJECTED y permiso de auditoría para CREATOR provisional. El API no sincroniza tablas automáticamente ni ejecuta DDL al arrancar.')
    paragraph(doc, 'El PDF describe siete etapas; el constructor conserva ocho porque Portada es una etapa independiente entre historia y planificación. Tipo, información general, historia/impacto, portada, plan/presupuesto, financiamiento, recompensas y revisión cubren el contenido solicitado. En donación se omite recompensas. La diferencia de presentación no elimina datos ni crea otra funcionalidad financiera.')

    for num, title, ids in [
        (3, 'Planificación presupuesto y riesgos', ['S2-02', 'S2-03']),
        (4, 'Financiamiento recompensas y preventas', ['S2-04', 'S2-05']),
        (5, 'Revisión final y envío de campaña', ['S2-10', 'S2-11']),
        (6, 'Decisiones administrativas y correcciones', ['S2-12', 'S2-13']),
        (7, 'Publicación estados y proyectos propios', ['S2-14', 'S2-15']),
        (8, 'Consulta pública y cola administrativa', ['S2-16', 'S2-17'])]:
        page(doc, f'{num} {title}')
        for ident in ids: task(doc, ident)

    page(doc, '9 Verificación documentación y calidad')
    heading(doc, 'S2 21 Verificar y documentar el segundo incremento', 2)
    paragraph(doc, 'Responsable planificado: Ricardo. Comprende las otras doce tareas y las regresiones de acceso de Sprint 1. El resultado incluye integración de las ramas del equipo, contratos, README, migraciones/permisos, guías y comprobaciones reproducibles. La evidencia publicada se puede localizar en DEV, commit 1fcecd4; no se requiere un PR para afirmar esta entrega porque se publicó directamente en la rama autorizada.')
    paragraph(doc, 'Criterios técnicos cubiertos: recorrido de campaña, datos reales, roles y propiedad, campos, fechas, importes, solicitudes duplicadas, persistencia, respuestas inciertas, instalación y documentación. Se conservan ensayos de escritorio, móvil y teclado del 5 de octubre. La aceptación final corresponde a una revisión externa y no se marca como concedida en este informe.')
    table(doc, ['Conjunto del último ensayo registrado', 'Aprobadas', 'Fallos', 'Omitidas'], [
        ['Frontend', '142', '0', '0'], ['Backend unitarias y HTTP', '131', '0', '0'],
        ['Procedimiento ADMIN local', '8', '0', '0'], ['Integración PostgreSQL real', '19', '0', '0'],
        ['Total', '300', '0', '0']], [3.7, 1.05, 1.0, 1.05])
    paragraph(doc, 'Este conteo corresponde al 5 de octubre, no a pruebas nuevas ejecutadas el día de preparación del informe. El corte previo de instalación limpia aprobó 276 casos; el de permisos, 291; el ajuste de navegación añadió nueve casos frontend, alcanzando 300. Los cortes son acumulativos y no deben sumarse entre sí.')
    paragraph(doc, 'Las integraciones crean fixtures efímeros y limpian sus propios registros. Comprueban el servidor real y PostgreSQL; el último ensayo utilizó DB_TEST_RESTART=true y confirmó conservación tras reiniciar el contenedor. Lint, tipos y compilación aprobaron. /api/health/ready respondió 200 con database=connected tanto directamente como a través de Vite.')
    paragraph(doc, 'En raíz: bun run check. En backend con pnpm 11.19.0: pnpm run check. Para integraciones locales: activar ALLOW_DB_TEST_WRITES=true y, si se quiere ensayar el reinicio, DB_TEST_RESTART=true; ejecutar pnpm run test:integration con la .env local preparada. Estas opciones nunca deben habilitarse en producción.')

    page(doc, '10 Validaciones seguridad y experiencia de usuario')
    paragraph(doc, 'La validación frontend reduce errores de interacción y localiza faltantes; la de backend protege el sistema aunque se envíe HTTP sin utilizar la interfaz. Los DTO rechazan campos extra y entradas anidadas inválidas. Las reglas de dominio verifican fechas reales, montos finitos positivos, precisión de dos decimales, duración entera y coherencia de modalidad. Los catálogos deben existir y estar activos.')
    paragraph(doc, 'Se aceptan Ñ, acentos y texto de campaña con números. Los nombres de persona no aceptan una secuencia puramente numérica. Las contraseñas se comparan exactamente, incluidos espacios y representación Unicode, sin recorte ni normalización. El límite se cuenta por puntos de código y el campo HTML no trunca incorrectamente los caracteres admitidos. NUL se rechaza antes de PostgreSQL con errores de campo, no se elimina silenciosamente.')
    paragraph(doc, 'La suma de presupuesto usa céntimos de los importes planificados validados; 0.10 y 0.20 producen 0.30. Cantidad, precio unitario e importe planificado son campos distintos: no se impuso una igualdad comercial no confirmada entre ellos, ni entre presupuesto y meta. El total no se presenta como un motor contable o de pagos.')
    paragraph(doc, 'Las contraseñas permanecen protegidas por scrypt y sal aleatoria, heredados del Sprint 1. La sesión utiliza una cookie HttpOnly, SameSite y atributos ajustados al entorno; el token se conserva como hash en el servidor y puede revocarse. No se coloca un token de sesión en localStorage. La protección de origen y X-Brotar-Request controla mutaciones; roles, propietario y estado se validan en servidor.')
    paragraph(doc, 'Durante un guardado se muestran carga y confirmación, se evita el doble envío y se conservan campos ante un fallo. Un JSON inválido o una respuesta de identidad/estado inesperada no se anuncia como éxito. Un HTTP 409 exige actualizar ronda o estado antes de repetir; un 422 presenta faltantes concretos.')
    paragraph(doc, 'El incremento final evita otro login o registro cuando /auth/me confirma sesión, adapta Mi cuenta/Mis borradores y conserva continuar para el borrador seleccionado. Una caída de API no se interpreta como sesión anónima. Estas redirecciones no conceden ADMIN ni CREATOR. Los tests de navegación y contrato se repitieron; el ensayo visual completo posterior al apagado no se repitió.')
    paragraph(doc, 'Recuperación sigue en modo de desarrollo local mientras no exista un correo institucional configurado: su enlace se conserva en el buzón privado indicado por el README. No se afirma que se envió email. Las pruebas verifican uso único, caducidad y revocación de sesiones después de restablecer.')

    page(doc, '11 Recorrido de demostración de inicio a fin')
    paragraph(doc, 'Preparación. Levantar V2, API y frontend según el README y comprobar readiness. Usar dos cuentas ficticias separadas: creador propietario con REGISTERED_USER y CREATOR, y ADMIN. Registrar una cuenta no concede esos privilegios. El administrador se provisiona por el procedimiento local; ADMIN puede conceder CREATOR desde la interfaz provisional con motivo y confirmación. No se comparten claves administrativas en este informe.')
    steps = [
        'Entrar como creador. Abrir Mis borradores y crear una campaña. Elegir donación para el recorrido mínimo o recompensa para demostrar su etapa adicional. Crear otro borrador confirma que el primero permanece accesible.',
        'Guardar información general con categoría y ubicación, historia con problema/solución/beneficiarios/resultados e indicador, y portada válida con texto alternativo. Recargar y seleccionar el mismo ID para comprobar persistencia.',
        'En Plan y presupuesto, registrar actividad con responsable y fechas, hito si corresponde, partida con categoría e importe positivo y riesgo con mitigación. Guardar cada bloque. Probar un campo inválido antes de corregirlo.',
        'Guardar meta, moneda activa, duración/modelo y fechas coherentes. Para recompensa/preventa, guardar oferta con mínimo y entrega. En donación confirmar que el paso se omite sin exigir una recompensa.',
        'Abrir Revisión y envío. Seguir los enlaces a faltantes, guardar correcciones y enviar. Solo la confirmación del servidor y estado IN_REVIEW prueban recepción; un botón pulsado no basta.',
        'Entrar como ADMIN en otra sesión. Abrir Administración de campañas, consultar el expediente y solicitar cambios con motivo y observación. Verificar la ronda y que no aparece en el catálogo público.',
        'Volver como creador. En Mis proyectos abrir Corregir campaña, leer observaciones, guardar y reenviar. Confirmar nueva ronda sin pérdida de anteriores. El administrador aprueba el expediente vigente.',
        'Comprobar que APPROVED sigue privada. ADMIN publica por la acción independiente y confirmada. Abrir Explorar sin sesión, buscar/filtrar y comparar portada, historia, meta e impacto con lo guardado.',
        'En campañas de ensayo distintas, demostrar rechazo con motivo y descarte confirmado de DRAFT propio. No usar la campaña publicada ni SQL directo para simular esas transiciones. Cerrar sesión y comprobar denegación de rutas privadas.'
    ]
    for i, item in enumerate(steps, 1): paragraph(doc, f'{i}. {item}')

    page(doc, '12 Decisiones pendientes defensa y conclusión')
    heading(doc, 'Decisiones provisionales', 2)
    paragraph(doc, 'D05 autoriza provisionalmente que ADMIN publique una campaña ya aprobada y con confirmación. CAMPAIGN_ADMIN_PUBLICATION_ENABLED activa esa operación sin conceder rol. D06 mantiene pendiente la provisión definitiva y relación de CREATOR con KYC/KYB; la habilitación local registra motivo, actor y auditoría sin verificar identidad ni cambiar contraseña. Correo real, textos legales y reglas de verificación continúan pendientes.')
    heading(doc, 'Límites del resultado', 2)
    paragraph(doc, 'Las pruebas aprobadas sostienen el cierre técnico de este alcance, no ausencia absoluta de defectos ni seguridad certificada. La concordancia con Figma se examinó parcialmente, no al cien por ciento de sus pantallas. Sigue el aviso de Vite por un paquete de aproximadamente 505 kB. El fallo recurrente de sockets de Docker Windows se recuperó conservando temporales, pero no se declaró resuelto definitivamente.')
    heading(doc, 'Preguntas para la defensa', 2)
    paragraph(doc, '¿Por qué aprobar no publica? Porque Fase 3 define APPROVED como estado privado y PUBLISHED como público; la publicación tiene autorización y confirmación independientes. ¿Por qué no basta ocultar botones? Una petición HTTP puede evitar la interfaz; servidor verifica sesión, rol, propiedad y estado. ¿Qué evita dos decisiones? Transacción, bloqueo de campaña/ronda y comparación de expectedRound. ¿Qué ocurre al reiniciar Docker? El volumen conserva la base; la API vuelve a conectarse sin restaurar SQL.')
    heading(doc, 'Conclusión', 2)
    paragraph(doc, 'El Sprint 2 entrega un recorrido integrado de campañas con persistencia y revisión real, manteniendo la base del Sprint 1. Las 13 tareas tienen implementación/evidencia técnica publicada en DEV. Las decisiones comerciales futuras y la aceptación de líderes no se sustituyen por el cierre local; se mantienen separadas para continuar el proyecto sin incorporar módulos expresamente excluidos.')
    heading(doc, 'Referencias de verificación', 2)
    for path in ['docs/requisitos/Brotar_Tarea_Fase3_Modulo_Campanas.pdf', 'docs/planificacion-sprints.json', 'docs/ajuste-sprint-2-fase-3-2026-09-30.md', 'docs/cierre-tecnico-sprint-2-2026-10-05.md', 'docs/entrega-dev-sprint-2-2026-10-05.md', 'docs/decisiones/ADR-005-publicacion-administrativa-provisional.md', 'README.md']:
        assert (ROOT / path).exists(); code(doc, path)
    paragraph(doc, 'Los documentos previos conservan su corte histórico. El README, secciones 16 y 17, y el commit 1fcecd4 actualizan publicación, permisos, varios borradores y los 300 casos del último ensayo registrado.')
    save(doc, 'Informe_Detallado_Sprint_2_Brotar.docx')

def principles():
    doc = new_doc('Principios y patrones del Sprint 2', 'Fundamentos de diseño ubicación de código y justificación técnica')
    page(doc, '1 Enfoque de programación y criterio de aplicación')
    paragraph(doc, 'El equipo aplicó principios donde resolvían problemas concretos del módulo de campañas. Se conservaron código y arquitectura adecuados; se refactorizó la coordinación de formularios y se corrigieron validaciones. La evaluación debe relacionar cada concepto con una unidad de código, la razón de usarla, una prueba y un límite de la afirmación. No se añadieron jerarquías ni patrones solo para cumplir una lista.')
    paragraph(doc, 'KISS, DRY y YAGNI son principios de simplicidad, reutilización y alcance. SOLID orienta responsabilidades y contratos. Repository y Adapter son patrones de diseño. IoC e inyección son técnicas de ensamblaje; las transacciones y bloqueos son mecanismos de consistencia. Presentar todos como un único tipo de programación ocultaría diferencias relevantes para la defensa.')
    table(doc, ['Enfoque', 'Aplicación en Sprint 2', 'Límite'], [
        ['Orientación a objetos', 'Casos de uso y repositorios con colaboradores', 'Sin jerarquías artificiales por rol'],
        ['Funciones puras', 'Validadores y cálculo de faltantes', 'No todo el sistema es funcional puro'],
        ['Declarativo', 'React describe la UI desde props y estado', 'Los efectos siguen controlados por hooks'],
        ['Asíncrono', 'Solicitudes HTTP y operaciones de base', 'async no equivale a atomicidad'],
        ['Capas y puertos', 'Dominio y aplicación separados de infraestructura', 'El adaptador sí conoce SQL/TypeORM']], [1.35, 3.4, 2.05])
    paragraph(doc, 'Recorrido de referencia: CampaignSectionsEditor usa useCampaignSections; el hook invoca sectionsClient y requestApi; SectionsController valida HTTP y delega a CampaignSections; este usa reglas de dominio y el puerto SectionsRepository; TypeormSectionsRepository persiste en PostgreSQL. La respuesta debe validar su contrato antes de confirmar el resultado.')
    ref(doc, 'backend/test/architecture.test.ts', "test('BG-58")
    paragraph(doc, 'La comprobación de imports detecta dependencias prohibidas de dominio/aplicación hacia framework o infraestructura. No demuestra por sí sola que todo el diseño sea perfecto; se complementa con pruebas de comportamiento e integración.')

    method_titles = ['2 KISS y reglas simples', '3 DRY y reutilización de conocimiento', '4 YAGNI y control del alcance', '5 SOLID y responsabilidad única', '6 SOLID y abierto cerrado', '7 SOLID y sustitución de Liskov', '8 SOLID y segregación de interfaces', '9 SOLID e inversión de dependencias', '10 IoC inyección y composición', '11 Boy Scout y mejoras acotadas']
    for i, method in enumerate(METHODS):
        title, definition, application, references, demo, limit, answer = method
        page(doc, method_titles[i])
        heading(doc, 'Concepto y problema que resuelve', 2)
        paragraph(doc, definition)
        heading(doc, 'Aplicación y justificación en Brotar', 2)
        paragraph(doc, application)
        if i == 0:
            code(doc, "isPublicCampaign(status) => status === 'PUBLISHED'")
            paragraph(doc, 'Esta función expresa la separación exigida entre aprobación y visibilidad pública. No se necesita un motor configurable de estados para comprobar esta regla acotada; los bloqueos y transacciones se mantienen en persistencia porque resuelven un problema diferente.')
        elif i == 1:
            paragraph(doc, 'El componente CreatorAccessActions también reutiliza los accesos superior e inferior de Para creadores. Ambos reciben la misma sesión confirmada desde PublicLayout, evitando decisiones visuales contradictorias. La protección real permanece en el servidor.')
        elif i == 3:
            paragraph(doc, 'La separación no convierte el componente en un archivo sin lógica: conserva decisiones de presentación y eventos. El hook concentra coordinación del estado y del guardado; el cliente valida el contrato de red. La responsabilidad se valora por razones de cambio, no por longitud del archivo.')
        elif i == 4:
            paragraph(doc, 'Un punto extensible ya utilizado es el parámetro send de requestApi. Las pruebas reemplazan el transporte sin reescribir el consumidor. La elección de fetch sigue siendo válida para este contrato; utilizar Axios por sí solo no demostraría abierto/cerrado ni mejor seguridad.')
        elif i == 7:
            paragraph(doc, 'DIP e inyección no son sinónimos. El puerto pertenece a la aplicación; NestJS decide qué implementación concreta entregar. Si el caso de uso recibiera una clase SQL concreta como dependencia de negocio, usar un contenedor no eliminaría ese acoplamiento.')
        heading(doc, 'Ubicación verificable en el corte DEV', 2)
        for r in references: ref(doc, *r)
        heading(doc, 'Demostración y evidencia', 2)
        paragraph(doc, demo)
        heading(doc, 'Alcance de la afirmación', 2)
        # Keep actual technical limits without referring to authoring tools.
        limit = limit.replace('Una prueba interactiva adicional del hook no fue repetida por bloqueo del navegador.', 'La cobertura automatizada del hook es acotada; no prueba por sí sola todas las interacciones posibles de la interfaz.')
        paragraph(doc, limit)
        heading(doc, 'Justificación para la defensa', 2)
        paragraph(doc, answer.replace('“', '').replace('”', '').replace('Separé ', 'Separamos '))

    page(doc, '12 Repository Adapter y límites de los patrones')
    heading(doc, 'Repository', 2)
    paragraph(doc, 'AdministrativeRepository, SubmissionRepository y PublicCampaignRepository describen operaciones que los casos de uso necesitan sin incluir SQL. TypeormAdministrativeRepository, TypeormSubmissionRepository y TypeormPublicCampaignRepository implementan esas operaciones sobre la base oficial. La separación permite probar coordinación con dobles y concentrar consultas, transacciones y mapeo en infraestructura.')
    ref(doc, 'backend/src/campaigns/application/administrative-review.ts', 'export interface AdministrativeRepository')
    ref(doc, 'backend/src/campaigns/infrastructure/typeorm-administrative.repository.ts', 'export class TypeormAdministrativeRepository')
    paragraph(doc, 'Ejemplo: AdministrativeReview solicita decide(actorId, id, input); no calcula nombres de tablas ni abre conexiones. El adaptador bloquea registros y persiste la decisión. Un doble de prueba puede comprobar que una entrada inválida no llega a escribir; PostgreSQL real debe comprobar la concurrencia que el doble no reproduce.')
    heading(doc, 'Adapter', 2)
    paragraph(doc, 'mapPublicCampaign traduce el contrato HTTP de catálogo al modelo Project utilizado por las tarjetas y el detalle. Verifica identificadores, importes, contenido e imagen. Los repositorios también adaptan los puertos de aplicación a PostgreSQL. El beneficio es mantener formatos de frontera fuera de la vista o del caso de uso, no ocultar un fallo con datos simulados.')
    ref(doc, 'src/features/public/campaignsApi.ts', 'export function mapPublicCampaign')
    paragraph(doc, 'Demostración: un contrato público malformado o una respuesta 404 no genera una tarjeta de ejemplo. Las pruebas de public-campaigns y administrative-campaigns comprueban que la lectura no anuncia éxito con datos anidados inválidos. La función mapPublicCampaign no concede acceso a una campaña privada: esa exclusión pertenece al backend.')
    heading(doc, 'Patrones no afirmados', 2)
    paragraph(doc, 'No se presenta Event Sourcing, CQRS completo, una máquina de estados genérica, Unit of Work propio ni una jerarquía GoF Strategy por cada modalidad. Tener tablas de historial no constituye Event Sourcing; separar lectura de escritura no prueba CQRS. El proyecto utiliza transacciones de TypeORM y reglas explícitas, suficientes para el alcance actual.')

    page(doc, '13 Asincronía transacciones y control de concurrencia')
    paragraph(doc, 'Promise.all carga secciones independientes del constructor. AbortController cancela lecturas obsoletas al abandonar una página; las comprobaciones de montaje evitan aplicar resultados a una vista ya cerrada. async/await mantiene legible la secuencia de guardado. Estas técnicas coordinan efectos, pero no garantizan que varias escrituras sean atómicas.')
    heading(doc, 'Lectura coherente del expediente', 2)
    paragraph(doc, 'El detalle administrativo se construye dentro de REPEATABLE READ usando el mismo EntityManager con consultas secuenciales. Así se evita presentar una mezcla de secciones correspondientes a snapshots distintos. La decisión se toma después sobre una ronda esperada; leer una revisión no autoriza a modificar cualquier versión futura.')
    ref(doc, 'backend/src/campaigns/infrastructure/typeorm-administrative.repository.ts', "transaction('REPEATABLE READ'")
    heading(doc, 'Escritura atómica y conflicto explícito', 2)
    paragraph(doc, 'La decisión y la publicación bloquean campaña y última ronda mediante SELECT FOR UPDATE. expectedRound verifica que la vista corresponda a la revisión vigente. Dentro de la transacción se validan estado/completitud y se actualizan revisión, campaña, observaciones e historial. Una falla provoca rollback; una segunda acción concurrente obtiene conflicto después de la transición válida.')
    ref(doc, 'backend/src/campaigns/infrastructure/typeorm-administrative.repository.ts', 'private async lock')
    code(doc, 'Bloquear campaña y ronda → validar estado y expectedRound → persistir → confirmar transacción')
    paragraph(doc, 'Se reutiliza el trigger oficial del historial y se añade motivo/contexto a su entrada, en lugar de producir un historial paralelo duplicado. APPROVED permanece privada; publicar exige una operación distinta y autorizada. La provisión de roles no forma parte de esta transacción de campaña.')
    heading(doc, 'Prueba que demuestra el beneficio', 2)
    ref(doc, 'backend/integration/campaign-sections-submission.integration.test.ts')
    paragraph(doc, 'La integración ejecuta envío, cambios, reenvío, aprobación, publicación y rechazo mediante APIs reales. Incluye entradas incompletas, catálogos deshabilitados y solicitudes concurrentes. Que un botón se deshabilite no sustituye esta protección: dos clientes independientes pueden enviar a la vez.')
    paragraph(doc, 'Los bloqueos y el aislamiento tienen coste y deben aplicarse a los datos relevantes. No se afirma serialización global del sistema ni garantía de disponibilidad bajo toda carga. La consistencia verificada corresponde a las transiciones y operaciones cubiertas por este incremento.')

    page(doc, '14 Validación defensa técnica y conclusión')
    paragraph(doc, 'La calidad del diseño se contrasta con comportamiento, no con la cantidad de principios nombrados. Los validadores del navegador mejoran la corrección del formulario; DTO y dominio vuelven a validar en servidor. La duplicación deliberada de reglas entre fronteras exige pruebas de paridad y no se resuelve eliminando una de las defensas.')
    table(doc, ['Pregunta del docente', 'Código para mostrar', 'Razón concreta'], [
        ['Dónde se separó lógica y vista', 'useCampaignSections y editor', 'Cambiar coordinación sin rehacer marcado'],
        ['Cómo se evita repetir HTTP', 'requestApi', 'Cookies, timeout y errores uniformes'],
        ['Cómo se prueba sin PostgreSQL', 'AdministrativeRepository y doble', 'Coordinar mediante un puerto'],
        ['Qué evita dos decisiones', 'lock y expectedRound', 'Estado y ronda comprobados bajo transacción'],
        ['Por qué no se agregó checkout', 'Alcance Fase 3 sección 11', 'Funcionalidad expresamente excluida']], [2.15, 2.15, 2.5])
    paragraph(doc, 'La refactorización extrajo el hook y corrigió decimales, fechas imposibles, NUL y límites Unicode. Las capas, puertos, hash y controles de sesión existentes se conservaron. El último ensayo registrado aprobó 300 casos, incluidos Sprint 1, procedimiento ADMIN e integración real; no se ejecutaron nuevamente para redactar este documento.')
    paragraph(doc, 'Ejemplos adicionales del incremento final son GuestAccessGate y CreatorAccessActions: cada uno concentra una decisión de acceso o presentación, comparte la sesión confirmada y evita un login redundante. El servidor continúa siendo autoridad. La recuperación solo notifica cambio de sesión tras confirmación, evitando convertir un fallo en una actualización engañosa.')
    ref(doc, 'src/features/access/session/GuestAccessGate.tsx', 'export function GuestAccessGate')
    ref(doc, 'tests/session-navigation.test.ts')
    paragraph(doc, 'Conclusión. Los principios se aplican de forma acotada y justificable al flujo de campañas. Separación de responsabilidades, puertos, adaptación de contratos y consistencia transaccional permiten defender por qué el código tiene esta estructura. Ningún principio sustituye autorización, pruebas de integración o aceptación final del cliente.')
    heading(doc, 'Referencias', 2)
    for path in ['README.md', 'backend/test/architecture.test.ts', 'tests/api-request.test.ts', 'tests/sprint2-quality.test.tsx', 'backend/test/sprint2-quality.test.ts', 'docs/cierre-tecnico-sprint-2-2026-10-05.md']:
        assert (ROOT / path).exists(); code(doc, path)
    paragraph(doc, 'Las rutas y líneas de este informe corresponden al commit 1fcecd4 de DEV. Al modificar el código posteriormente deben localizarse de nuevo por el símbolo indicado. Los reportes históricos registran comprobaciones anteriores, no decisiones finales de cliente.')
    save(doc, 'Principios_y_Patrones_Sprint_2_Brotar.docx')

if __name__ == '__main__':
    assert len(TASKS) == 13
    report()
    principles()
